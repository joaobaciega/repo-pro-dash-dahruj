"""Regressões da fronteira pública/restrita e da movimentação de estoque."""

import io
from pathlib import Path
import time
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

import estoque_pagina
from test_verbas import entrada, saida
from verbas import calcular_verbas, movimentacao_por_unidade
from verbas_pagina import COLUNAS_CLIENTE, exportar_excel_cliente, tabela_cliente


class ApresentacaoTest(unittest.TestCase):
    def test_excel_publico_nao_contem_detalhes_internos(self):
        m = calcular_verbas(pd.DataFrame([entrada()]), pd.DataFrame([saida()]), {}, {})
        tabela = tabela_cliente(m, 2026)
        # Mesmo que o chamador envie dados a mais, a exportação usa lista explícita.
        tabela["Estoque por unidade"] = "interno"
        wb = load_workbook(io.BytesIO(exportar_excel_cliente(tabela)), data_only=True)
        self.assertEqual(wb.sheetnames, ["Verbas mensais"])
        ws = wb.active
        self.assertEqual([c.value for c in ws[1]][1:], list(COLUNAS_CLIENTE.values()))
        self.assertEqual(ws.cell(ws.max_row, ws.max_column).value, 1450)
        self.assertTrue(all(c.data_type == "n" for row in ws.iter_rows(min_row=2, min_col=2) for c in row))

    def test_movimentacao_carrega_ano_anterior_e_preserva_negativos(self):
        v = pd.DataFrame([entrada("2025-12-01", qtde=100), entrada("2026-02-01", qtde=20)])
        s = pd.DataFrame([saida("2026-01-01", diant=40), saida("2026-02-01", diant=90)])
        m = movimentacao_por_unidade(v, s, "2026-02-01")
        row = m.loc["FIAT CEASA"]
        self.assertEqual(tuple(row[["inicial_diant", "entrada_diant", "saida_diant", "estoque_diant"]]), (60, 20, 90, -10))

    def test_estoque_bloqueado_nao_consulta_dados(self):
        with patch.object(estoque_pagina.st, "session_state", {}), \
             patch.object(estoque_pagina.st, "error") as erro, \
             patch.object(estoque_pagina.db, "ler_vendas_verbas") as consulta:
            estoque_pagina.pagina_estoque()
            erro.assert_called_once()
            consulta.assert_not_called()


# Usa a navegação real, isolando as outras páginas e o banco. O arquivo fica
# no cache local porque diretórios temporários do Windows podem negar escrita.
DRIVER = '''
from pathlib import Path
import pandas as pd
import streamlit as st
import db
from test_verbas import entrada, saida

source = Path(__file__).resolve().parents[1] / "app.py"
code = source.read_text(encoding="utf-8")
offset = code.index("# ============================ NAVEGAÇÃO")
scope = {"__name__": "__main__", "__file__": str(source)}
exec(compile(code[:offset], str(source), "exec"), scope)
v = pd.DataFrame([entrada(), entrada(qtde=10, cliente="CLIENTE - NISSAN B.LEME")])
s = pd.DataFrame([saida(), saida(diant=12, unidade="Nissan Braz Leme")])
db.ler_vendas_verbas = lambda: v
db.ler_base_tidy = lambda: s
db.ler_pagamentos_verba = lambda: {pd.Timestamp("2026-02-01"): {"consultor": True, "gerente": True}}
db.ler_pagamentos_marketing = lambda: {}
scope["pagina_dashboard"] = lambda: st.write("Dashboard")
scope["pagina_lancamento"] = lambda: st.write("Lançamento")
scope["pagina_historico"] = lambda: st.write("Histórico")
scope["_senha_esperada"] = lambda: "senha-de-teste"
exec(compile(code[offset:], str(source), "exec"), scope)
'''


class NavegacaoTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from streamlit.testing.v1 import AppTest
        cls.AppTest = AppTest
        cls.driver = Path(__file__).resolve().parent / "__pycache__" / "verbas_pages_test.py"
        cls.driver.parent.mkdir(exist_ok=True)
        cls.driver.write_text(DRIVER, encoding="utf-8")

    @classmethod
    def tearDownClass(cls):
        cls.driver.unlink(missing_ok=True)

    def app(self, restrito=False):
        app = self.AppTest.from_file(str(self.driver))
        app.session_state["menu_pagina"] = "Verbas"
        if restrito:
            app.session_state["lanc_ok"] = True
            app.session_state["lanc_visto"] = time.monotonic()
        return app.run(timeout=20)

    def test_publico_tem_tres_cards_sem_avisos_ou_estoque(self):
        app = self.app()
        self.assertFalse(app.exception)
        self.assertEqual([m.label for m in app.metric], ["Verba gerada", "Verba paga", "Saldo de verbas"])
        self.assertFalse(app.warning)
        self.assertFalse(app.expander)
        self.assertNotIn("📦 Estoque", app.sidebar.radio[0].options)
        self.assertFalse(any("estoque" in str(df.value.columns).lower() for df in app.dataframe))
        app.selectbox[1].select("Fevereiro").run()
        self.assertFalse(app.exception)
        self.assertEqual(app.metric[0].value, "R$ 0,00")
        self.assertEqual(app.metric[1].value, "R$ 650,00")
        self.assertEqual(app.metric[2].value, "R$ 945,00")

    def test_estoque_so_com_acesso_e_filtro_unidade(self):
        app = self.app(restrito=True)
        self.assertIn("📦 Estoque", app.sidebar.radio[0].options)
        app.sidebar.radio[0].set_value("📦 Estoque").run(timeout=20)
        self.assertFalse(app.exception)
        self.assertEqual(app.metric[0].value, "58")
        self.assertTrue(app.warning)
        self.assertEqual(app.metric[3].value, "1")
        self.assertTrue(any("Oferecer novo pedido" in i.value for i in app.info))
        app.multiselect[0].set_value(["FIAT CEASA"]).run(timeout=20)
        self.assertFalse(app.exception)
        self.assertEqual(app.metric[0].value, "60")
        self.assertFalse(app.warning)
        self.assertEqual(app.metric[3].value, "0")
        app.button(key="btn_cadeado").click().run(timeout=20)
        self.assertFalse(app.exception)
        self.assertNotIn("📦 Estoque", app.sidebar.radio[0].options)
        self.assertEqual(app.sidebar.radio[0].value, "Dashboard")
        self.assertFalse(app.metric)

    def test_sessao_expirada_retira_estoque(self):
        app = self.app(restrito=True)
        app.session_state["menu_pagina"] = "📦 Estoque"
        app.session_state["lanc_visto"] = time.monotonic() - 1801
        app.run(timeout=20)
        self.assertFalse(app.exception)
        self.assertNotIn("📦 Estoque", app.sidebar.radio[0].options)
        self.assertFalse(app.dataframe)

    def test_cadeado_exige_senha_e_libera_estoque(self):
        app = self.app()
        app.button(key="btn_cadeado").click().run(timeout=20)
        self.assertFalse(app.exception)
        app.text_input[0].set_value("incorreta")
        # AppTest refaz o script inteiro; no navegador o formulário refaz só
        # o fragmento do diálogo. Reabre o diálogo para processar a submissão.
        app.session_state["pedir_senha_restrita"] = True
        next(b for b in app.button if b.label == "Entrar").click().run(timeout=20)
        self.assertNotIn("📦 Estoque", app.sidebar.radio[0].options)
        self.assertTrue(app.error)
        app.text_input[0].set_value("senha-de-teste")
        app.session_state["pedir_senha_restrita"] = True
        next(b for b in app.button if b.label == "Entrar").click().run(timeout=20)
        self.assertFalse(app.exception)
        self.assertIn("📦 Estoque", app.sidebar.radio[0].options)


if __name__ == "__main__":
    unittest.main()
