import io
import unittest

import pandas as pd
from openpyxl import load_workbook

from estoque_pagina import exportar_estoque
from test_verbas import entrada, saida
from verbas import movimentacao_por_unidade, reposicao_por_unidade


class ReposicaoTest(unittest.TestCase):
    def test_media_mensal_soma_consultores_inclui_zero_e_anos_anteriores(self):
        v = pd.DataFrame([entrada("2025-11-01", qtde=70)])
        s = pd.DataFrame([saida("2025-11-01", diant=20), saida("2025-11-01", diant=10),
                          saida("2025-12-01", diant=0), saida("2026-02-01", diant=30),
                          saida("2026-03-01", diant=1000)])
        # Novembro soma 30, dezembro 0, fevereiro 30: 60 / 3 = 20.
        # Janeiro não tem registros e março fica fora da posição selecionada.
        r = reposicao_por_unidade(v, s, "2026-02-01").iloc[0]
        self.assertEqual(r["Meses com registro"], 3)
        self.assertEqual(r["Consumo médio/mês"], 20)
        self.assertEqual(r["Estoque"], 10)
        self.assertEqual(r["Cobertura (meses)"], 0.5)
        self.assertTrue(r["Oferecer pedido"])

    def test_tipos_independentes_e_limite_estrito(self):
        v = pd.DataFrame([entrada(qtde=100), entrada(qtde=19, tipo="tras")])
        s = pd.DataFrame([saida(diant=50, tras=10)])
        r = reposicao_por_unidade(v, s, "2026-02-01")
        self.assertFalse(r.iloc[0]["Oferecer pedido"])  # Estoque igual à média.
        self.assertTrue(r.iloc[1]["Oferecer pedido"])   # Traseiro não compensado pelo dianteiro.

    def test_sem_historico_sem_consumo_e_saldo_negativo(self):
        v = pd.DataFrame([entrada(qtde=10), entrada(qtde=10, cliente="CLIENTE - HONDA CEASA")])
        s = pd.DataFrame([saida(diant=12)])
        r = reposicao_por_unidade(v, s, "2026-02-01")
        fiat = r[r.Unidade == "FIAT CEASA"]
        honda = r[r.Unidade == "HONDA CEASA"]
        self.assertTrue(fiat.iloc[0]["Oferecer pedido"])
        self.assertIn("conferir saldo negativo", fiat.iloc[0]["Situação"])
        self.assertEqual(fiat.iloc[1]["Situação"], "Sem consumo registrado")
        self.assertTrue(pd.isna(fiat.iloc[1]["Cobertura (meses)"]))
        self.assertTrue((honda["Situação"] == "Sem histórico").all())
        self.assertFalse(honda["Oferecer pedido"].any())

    def test_alias_e_inicio_individual_do_historico(self):
        v = pd.DataFrame([entrada(qtde=100), entrada(qtde=10, cliente="CLIENTE - NISSAN B.LEME")])
        s = pd.DataFrame([saida("2026-01-01", diant=20), saida("2026-02-01", diant=20),
                          saida("2026-02-01", diant=6, unidade="Nissan Braz Leme")])
        r = reposicao_por_unidade(v, s, "2026-02-01")
        nissan = r[(r.Unidade == "NISSAN BRAZ LEME") & (r.Tipo == "Dianteiros (pares)")].iloc[0]
        self.assertEqual(nissan["Meses com registro"], 1)
        self.assertEqual(nissan["Consumo médio/mês"], 6)
        self.assertTrue(nissan["Oferecer pedido"])

    def test_exportacao_preserva_media_fracionaria_e_situacao(self):
        v = pd.DataFrame([entrada(qtde=15)])
        s = pd.DataFrame([saida("2026-01-01", diant=5), saida("2026-02-01", diant=6)])
        r = reposicao_por_unidade(v, s, "2026-02-01")
        movimento = movimentacao_por_unidade(v, s, "2026-02-01")
        wb = load_workbook(io.BytesIO(exportar_estoque(movimento, movimento, pd.DataFrame(), r)), data_only=True)
        ws = wb["Reposição"]
        cols = {c.value: c.column for c in ws[1]}
        self.assertEqual(ws.cell(2, cols["Consumo médio/mês"]).value, 5.5)
        self.assertEqual(ws.cell(2, cols["Consumo médio/mês"]).number_format, "0.00")
        self.assertEqual(ws.cell(2, cols["Situação"]).value, "Oferecer novo pedido")


if __name__ == "__main__":
    unittest.main()
