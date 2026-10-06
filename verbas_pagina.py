"""Aba Verbas: provisão, vendas ao consumidor, estoque e pagamentos."""

import io

import pandas as pd
import plotly.express as px
import streamlit as st

import db
from verbas import CATEGORIAS, calcular_verbas


MESES = ["", "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho",
         "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]
COLUNAS = {
    "faturamento": "Faturamento Sell In",
    "entrada_diant": "Sell In dianteiros (pares)", "entrada_tras": "Sell In traseiros (un)",
    "saida_diant": "Sell Out dianteiros (pares)", "saida_tras": "Sell Out traseiros (un)",
    "estoque_diant": "Estoque dianteiros (pares)", "estoque_tras": "Estoque traseiros (un)",
    "gerada_consultor": "Provisão Consultor", "gerada_gerente": "Provisão Gerente",
    "gerada_marketing": "Provisão Marketing", "gerada_total": "Total provisionado (Sell In)",
    "realizada_consultor": "Devido Consultor (Sell Out)",
    "realizada_gerente": "Devido Gerente (Sell Out)",
    "realizada_marketing": "Marketing vinculado ao Sell Out",
    "paga_consultor": "Consultor Pago", "paga_gerente": "Gerente Pago",
    "paga_marketing": "Marketing Pago", "paga_total": "Total Pago",
    "estoque_consultor": "Verba em estoque Consultor", "estoque_gerente": "Verba em estoque Gerente",
    "estoque_marketing": "Verba em estoque Marketing", "estoque_total": "Total de verba em estoque",
    "a_pagar_consultor": "A pagar Consultor no mês", "a_pagar_gerente": "A pagar Gerente no mês",
    "pendente_consultor": "A pagar Consultor acumulado", "pendente_gerente": "A pagar Gerente acumulado",
    "saldo_consultor": "Variação saldo Consultor", "saldo_gerente": "Variação saldo Gerente",
    "saldo_marketing": "Variação saldo Marketing", "saldo_total": "Variação saldo total",
    "acumulado_consultor": "Saldo acumulado Consultor", "acumulado_gerente": "Saldo acumulado Gerente",
    "acumulado_marketing": "Saldo acumulado Marketing", "acumulado_total": "Saldo acumulado total",
}
QUANTIDADES = {c for c in COLUNAS if c.endswith(("_diant", "_tras"))}
SALDOS = {c for c in COLUNAS if c.startswith(("estoque_", "acumulado_", "pendente_"))}


def dinheiro(valor):
    return (f"R$ {valor:,.2f}").replace(",", "X").replace(".", ",").replace("X", ".")


def tabela_ano(mensal, ano, pagos):
    """Fluxos são somados; estoque/saldo na linha total são posições finais."""
    rec = mensal[mensal.index.year == ano]
    tabela = rec[list(COLUNAS)].copy()
    tabela.index = [MESES[m.month] for m in rec.index]
    tabela["status"] = [
        "Sem Sell Out" if not (r.saida_diant or r.saida_tras) else
        "Pago" if all(pagos.get(m, {}).get(c) for c in ("consultor", "gerente")) else
        "Parcial" if any(pagos.get(m, {}).get(c) for c in ("consultor", "gerente")) else "Em aberto"
        for m, r in rec.iterrows()
    ]
    if not rec.empty:
        tabela.loc["TOTAL / POSIÇÃO FINAL"] = {
            **{c: rec[c].iloc[-1] if c in SALDOS else rec[c].sum() for c in COLUNAS}, "status": ""
        }
    return tabela.rename(columns={**COLUNAS, "status": "Status do Sell Out"}).rename_axis("Mês")


def exportar_excel(tabela, estoque):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        tabela.to_excel(writer, sheet_name="Verbas mensais")
        estoque.to_excel(writer, sheet_name="Estoque por unidade")
        pd.DataFrame({"Metodologia": [
            "Sell In provisiona a verba. Sell Out gera o valor devido a consultores e gerentes.",
            "Pago = valor devido pelo Sell Out do mês, apenas quando marcado como pago na planilha.",
            "Estoque = Sell In acumulado menos Sell Out acumulado, desde o início das bases (estoque inicial zero).",
            "Marketing mantém a reserva do Sell In, descontando os pagamentos efetivamente lançados.",
            "Verba de marketing em estoque é informativa e já integra o saldo de marketing; não somar novamente.",
            "TOTAL / POSIÇÃO FINAL soma fluxos; estoques, valores pendentes e saldos acumulados usam a última posição.",
            "Tarifas por tipo vêm da base de Sell In. Divergências de tarifa impedem o cálculo automático.",
            "Valores pagos a consultores/gerentes são calculados das marcações mensais, não de comprovantes bancários.",
        ]}).to_excel(writer, sheet_name="Metodologia", index=False)
        from openpyxl.styles import Font, PatternFill
        for ws in writer.book:
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions
            for cel in ws[1]:
                cel.font = Font(bold=True, color="FFFFFF")
                cel.fill = PatternFill("solid", fgColor="EB5E33")
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(50, max(18, len(str(col[0].value)) + 2))
            if ws.title == "Verbas mensais":
                for col in ws.iter_cols(min_col=2):
                    formato = "#,##0" if col[0].value in {COLUNAS[c] for c in QUANTIDADES} else 'R$ #,##0.00'
                    for cel in col[1:]:
                        if isinstance(cel.value, (int, float)):
                            cel.number_format = formato
    return buf.getvalue()


# Lista explícita: o relatório público nunca recebe os campos de estoque ou
# as planilhas de conferência usadas na área interna.
COLUNAS_CLIENTE = {
    "gerada_consultor": "Gerado — Consultor",
    "gerada_gerente": "Gerado — Gerente",
    "gerada_marketing": "Gerado — Marketing",
    "gerada_total": "Total gerado",
    "paga_consultor": "Pago — Consultor",
    "paga_gerente": "Pago — Gerente",
    "paga_marketing": "Pago — Marketing",
    "paga_total": "Total pago",
    "acumulado_total": "Saldo de verbas",
}


def tabela_cliente(mensal, ano):
    rec = mensal.loc[mensal.index.year == ano, list(COLUNAS_CLIENTE)].copy()
    tabela = rec.copy()
    tabela.index = [MESES[m.month] for m in rec.index]
    if not rec.empty:
        tabela.loc["TOTAL DO ANO / SALDO FINAL"] = {
            c: rec[c].iloc[-1] if c == "acumulado_total" else rec[c].sum()
            for c in COLUNAS_CLIENTE
        }
    return tabela.rename(columns=COLUNAS_CLIENTE).rename_axis("Mês")


def exportar_excel_cliente(tabela):
    """Exporta exclusivamente os mesmos dados públicos da tabela mensal."""
    tabela = tabela[list(COLUNAS_CLIENTE.values())]
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        tabela.to_excel(writer, sheet_name="Verbas mensais")
        ws = writer.book["Verbas mensais"]
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        from openpyxl.styles import Font, PatternFill
        for cel in ws[1]:
            cel.font = Font(bold=True, color="FFFFFF")
            cel.fill = PatternFill("solid", fgColor="EB5E33")
        ws.column_dimensions["A"].width = 34
        for col in ws.iter_cols(min_col=2):
            ws.column_dimensions[col[0].column_letter].width = 23
            for cel in col[1:]:
                cel.number_format = 'R$ #,##0.00'
    return buf.getvalue()


def pagina_verbas():
    st.title("💰 Verbas")
    try:
        vendas = db.ler_vendas_verbas()
        sell_out = db.ler_base_tidy()
        pagos = db.ler_pagamentos_verba()
        mkt_pagos = db.ler_pagamentos_marketing()
        if vendas.empty or sell_out.empty:
            st.info("As informações de verbas ainda não estão disponíveis.")
            return
        mensal = calcular_verbas(vendas, sell_out, pagos, mkt_pagos)
    except Exception:
        st.error("As informações de verbas estão temporariamente indisponíveis. Tente novamente mais tarde.")
        return

    ca, cm, _ = st.columns([1, 2, 3])
    ano = ca.selectbox("Ano", sorted(set(mensal.index.year), reverse=True))
    do_ano = mensal[mensal.index.year == ano]
    escolha = cm.selectbox("📅 Mês", ["Ano todo"] + [MESES[m.month] for m in reversed(do_ano.index)])
    rec = do_ano if escolha == "Ano todo" else do_ano[do_ano.index.month == MESES.index(escolha)]
    fim = rec.index[-1]
    periodo = str(ano) if escolha == "Ano todo" else f"{escolha}/{ano}"
    fechamento = f"{MESES[fim.month]}/{fim.year}"
    soma, posicao = rec.sum(numeric_only=True), rec.iloc[-1]
    st.caption(f"Período: {periodo} · Valores em reais. O mês mais recente pode estar em andamento.")

    with st.container(key="verba_gerada"):
        cards = st.columns(3)
        cards[0].metric("Verba gerada", dinheiro(soma.gerada_total),
                        help="Verba gerada pelas compras no período selecionado.")
        cards[1].metric("Verba paga", dinheiro(soma.paga_total),
                        help="Pagamentos de consultores, gerentes e marketing no período selecionado.")
        cards[2].metric("Saldo de verbas", dinheiro(posicao.acumulado_total),
                        help=f"Total gerado menos pagamentos até {fechamento}, incluindo os períodos anteriores. Não representa apenas valores já devidos para pagamento.")
    st.caption(f"Gerado e pago referem-se ao período selecionado. O saldo considera todo o histórico até {fechamento}.")

    resumo = pd.DataFrame([
        {"Categoria": cat.title(), "Gerado no período": soma[f"gerada_{cat}"],
         "Pago no período": soma[f"paga_{cat}"], "Saldo de verbas": posicao[f"acumulado_{cat}"]}
        for cat in CATEGORIAS
    ])
    st.dataframe(resumo.style.format({c: dinheiro for c in resumo.columns if c != "Categoria"}),
                 width="stretch", hide_index=True)
    st.caption("Os pagamentos a consultores e gerentes consideram as vendas ao consumidor. Marketing considera os valores utilizados nas ações realizadas.")

    st.subheader(f"Verbas mês a mês — {ano}")
    tabela = tabela_cliente(mensal, ano)
    st.dataframe(tabela.style.format(dinheiro), width="stretch")
    st.caption("O saldo de cada mês inclui os períodos anteriores. Na última linha, gerado e pago são os totais do ano; o saldo é a posição ao final do período.")
    st.download_button("Exportar relatório (.xlsx)", exportar_excel_cliente(tabela),
                       file_name=f"verbas_{ano}.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    st.subheader("Verba gerada e paga")
    grafico = do_ano[["gerada_total", "paga_total"]].rename(columns={
        "gerada_total": "Gerada", "paga_total": "Paga"})
    grafico.index = [MESES[m.month][:3] for m in grafico.index]
    fig = px.bar(grafico, barmode="group", color_discrete_sequence=["#EB5E33", "#F5A623"])
    fig.update_layout(template="plotly_dark", xaxis_title=None, yaxis_title=None,
                      legend_title=None, paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)")
    fig.update_yaxes(tickprefix="R$ ")
    st.plotly_chart(fig, width="stretch")


