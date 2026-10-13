import io
import unittest

import pandas as pd
from openpyxl import load_workbook

from verbas import calcular_verbas, estoque_por_unidade
from verbas_pagina import exportar_excel, tabela_ano, COLUNAS


def entrada(data="2026-01-15", qtde=100, tipo="diant", cliente="CLIENTE - FIAT CEASA"):
    fator = 1 if tipo == "diant" else 0.5
    return dict(data=pd.Timestamp(data), qtde=qtde, tipo_refil=tipo, cliente=cliente,
                total_item=qtde * 84.9, verba_consultor=10 * fator, verba_gerente=2.5 * fator,
                verba_reserva=2 * fator, total_consultor=qtde * 10 * fator,
                total_gerente=qtde * 2.5 * fator, total_reserva=qtde * 2 * fator)


def saida(mes="2026-02-01", diant=40, tras=0, unidade="Fiat Ceasa"):
    return dict(mes=pd.Timestamp(mes), refil_diant=diant, refil_tras=tras, unidade=unidade)


class VerbasTest(unittest.TestCase):
    def calcular(self, entradas=None, saidas=None, pagos=None, marketing=None):
        return calcular_verbas(pd.DataFrame(entradas or [entrada()]),
                               pd.DataFrame(saidas or [saida()]), pagos or {}, marketing or {})

    def test_compra_em_janeiro_venda_em_fevereiro(self):
        pagos = {pd.Timestamp("2026-01-01"): {"consultor": True, "gerente": True},
                 pd.Timestamp("2026-02-01"): {"consultor": True}}
        m = self.calcular(pagos=pagos)
        jan, fev = m.iloc[0], m.iloc[1]
        self.assertEqual(jan.paga_total, 0)
        self.assertEqual(fev.paga_consultor, 400)
        self.assertEqual(fev.paga_gerente, 0)
        self.assertEqual(fev.estoque_total, 870)
        self.assertEqual(fev.pendente_gerente, 100)
        self.assertEqual(fev.acumulado_consultor, 600)
        self.assertEqual(fev.acumulado_marketing, 200)

    def test_estoque_nao_reinicia_no_ano_e_total_nao_soma_posicoes(self):
        m = self.calcular(entradas=[entrada("2025-12-01")], saidas=[saida("2026-02-01"), saida("2026-03-01", 10)])
        tabela = tabela_ano(m, 2026, {})
        self.assertEqual(m.loc["2026-01-01", "estoque_diant"], 100)
        self.assertEqual(tabela.iloc[-1][COLUNAS["estoque_diant"]], 50)
        self.assertEqual(tabela.iloc[-1][COLUNAS["saida_diant"]], 50)
        self.assertEqual(tabela.iloc[-1][COLUNAS["acumulado_total"]], 1450)

    def test_marketing_em_mes_sem_vendas_e_saldo_negativo(self):
        m = self.calcular(marketing={pd.Timestamp("2026-03-01"): 300})
        self.assertEqual(m.iloc[-1].paga_marketing, 300)
        self.assertEqual(m.iloc[-1].acumulado_marketing, -100)
        self.assertEqual(m.iloc[-1].estoque_diant, 60)

    def test_dianteiros_e_traseiros_sem_converter_par_em_duas_unidades(self):
        m = self.calcular(entradas=[entrada(), entrada(tipo="tras", qtde=50)], saidas=[saida(diant=40, tras=20)])
        self.assertEqual(m.iloc[-1].realizada_consultor, 500)
        self.assertEqual(m.iloc[-1].realizada_gerente, 125)
        self.assertEqual(m.iloc[-1].estoque_tras, 30)

    def test_atualizacao_sell_out_recalcula_pagamento(self):
        pagos = {pd.Timestamp("2026-02-01"): {"consultor": True, "gerente": True}}
        antes = self.calcular(pagos=pagos)
        depois = self.calcular(saidas=[saida(diant=50)], pagos=pagos)
        self.assertEqual(depois.iloc[-1].paga_total - antes.iloc[-1].paga_total, 125)
        self.assertEqual(depois.iloc[-1].estoque_diant, 50)

    def test_diferencas_negativas_visiveis(self):
        m = self.calcular(saidas=[saida(diant=110)])
        self.assertEqual(m.iloc[-1].estoque_diant, -10)
        self.assertEqual(m.iloc[-1].realizada_consultor, 1100)

    def test_nissan_consolidada_e_alias_interlagos(self):
        v = pd.DataFrame([entrada(qtde=190, cliente="CLIENTE - NISSAN B.LEME"),
                          entrada(qtde=940, cliente="CLIENTE - NISSAN CEASA"),
                          entrada(qtde=500, cliente="CLIENTE - JEEP INTERLAGOS")])
        s = pd.DataFrame([saida(diant=131, unidade="Nissan Braz Leme"),
                          saida(diant=773, unidade="Nissan Ceasa"), saida(diant=470, unidade="Jeep W. Luiz")])
        e = estoque_por_unidade(v, s, "2026-02-01")
        self.assertEqual(e.iloc[:, 0].to_dict(), {"JEEP W LUIZ": 30, "NISSAN CEASA": 226})

    def test_falha_em_tipo_tarifa_ou_total_ambiguo(self):
        for invalido in [dict(tipo_refil=None), dict(verba_consultor=11), dict(total_consultor=999), dict(qtde=-1)]:
            with self.subTest(invalido=invalido), self.assertRaises(ValueError):
                self.calcular(entradas=[entrada(), {**entrada(), **invalido}])

    def test_exportacao_numerica_confere_com_tela(self):
        m = self.calcular()
        tabela = tabela_ano(m, 2026, {})
        estoque = estoque_por_unidade(pd.DataFrame([entrada()]), pd.DataFrame([saida()]), "2026-02-01")
        wb = load_workbook(io.BytesIO(exportar_excel(tabela, estoque)), data_only=True)
        ws = wb["Verbas mensais"]
        cab = {c.value: c.column for c in ws[1]}
        valor = ws.cell(ws.max_row, cab[COLUNAS["estoque_total"]])
        self.assertEqual(valor.value, 870)
        self.assertEqual(valor.data_type, "n")
        self.assertIn("R$", valor.number_format)


if __name__ == "__main__":
    unittest.main()
