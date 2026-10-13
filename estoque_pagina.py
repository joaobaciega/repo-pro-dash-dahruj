"""Acompanhamento interno de estoque, protegido pelo acesso da aplicação."""

import io
import math
from decimal import Decimal, ROUND_HALF_UP

import pandas as pd
import plotly.express as px
import streamlit as st

import db
from verbas import CATEGORIAS, calcular_verbas, movimentacao_por_unidade, reposicao_por_unidade
from verbas_pagina import MESES, dinheiro


COLUNAS_ESTOQUE = {
    "inicial_diant": "Inicial diant. (pares)", "entrada_diant": "Entradas diant. (pares)",
    "saida_diant": "Saídas diant. (pares)", "estoque_diant": "Saldo diant. (pares)",
    "inicial_tras": "Inicial tras. (un)", "entrada_tras": "Entradas tras. (un)",
    "saida_tras": "Saídas tras. (un)", "estoque_tras": "Saldo tras. (un)",
}

COLUNAS_REPOSICAO = ["Unidade", "Saldo dianteiro (pares)", "Consumo dianteiro/mês",
                    "Duração dianteiros", "Saldo traseiro (un)", "Consumo traseiro/mês",
                    "Duração traseiros", "Oferecer pedido de"]


def duracao_estoque(cobertura):
    """Meses de 30 dias, sem frações e sem arredondar a duração para cima."""
    if pd.isna(cobertura):
        return "—"
    if cobertura < 0:
        return "Conferir saldo"
    dias = math.floor(cobertura * 30 + 1e-9)
    if dias == 0 and cobertura > 0:
        return "<1d"
    if dias <= 30:
        return f"{dias}d"
    meses, resto = divmod(dias, 30)
    return f"{meses}m e {resto}d" if resto else f"{meses}m"


def tabela_reposicao(reposicao):
    """Uma linha por loja, com os dois produtos lado a lado. Só apresentação."""
    linhas = []
    for unidade, grupo in reposicao.groupby("Unidade", sort=True):
        linha = {"Unidade": unidade}
        pedidos = []
        for tipo, sufixo, saldo_col in (("Dianteiros (pares)", "dianteiro", "Saldo dianteiro (pares)"),
                                       ("Traseiros (un)", "traseiro", "Saldo traseiro (un)")):
            r = grupo[grupo["Tipo"] == tipo].iloc[0]
            media = r["Consumo médio/mês"]
            linha[saldo_col] = int(r["Estoque"])
            linha[f"Consumo {sufixo}/mês"] = (None if pd.isna(media) else
                int(Decimal(str(media)).quantize(Decimal("1"), rounding=ROUND_HALF_UP)))
            duracao = duracao_estoque(r["Cobertura (meses)"])
            if not r["Meses com registro"]:
                duracao = "Sem histórico"
            elif media == 0:
                duracao = "Sem consumo"
            if r["Estoque"] < 0:
                duracao = "Conferir saldo"
            linha[f"Duração {sufixo}s"] = duracao
            if r["Oferecer pedido"]:
                pedidos.append(sufixo + "s")
        linha["Oferecer pedido de"] = " e ".join(pedidos).capitalize() if pedidos else "—"
        linhas.append(linha)
    return pd.DataFrame(linhas, columns=COLUNAS_REPOSICAO)


def destacar_reposicao(linha):
    destaque = "background-color: #713D13; color: #FFFFFF; font-weight: bold"
    pedido = linha["Oferecer pedido de"].lower()
    return [destaque if (c in ("Unidade", "Oferecer pedido de") and pedido != "—")
            or ("dianteiro" in c.lower() and "dianteiros" in pedido)
            or ("traseiro" in c.lower() and "traseiros" in pedido) else ""
            for c in linha.index]


def exportar_estoque(movimento, historico, verba, reposicao=None):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        ordem = ["estoque_diant", "estoque_tras", "inicial_diant", "entrada_diant",
                 "saida_diant", "inicial_tras", "entrada_tras", "saida_tras"]
        movimento[ordem].rename(columns=COLUNAS_ESTOQUE).to_excel(writer, sheet_name="Estoque por unidade")
        historico.rename(columns=COLUNAS_ESTOQUE).to_excel(writer, sheet_name="Evolução mensal")
        verba.to_excel(writer, sheet_name="Verba em estoque", index=False)
        if reposicao is not None:
            tabela_reposicao(reposicao).to_excel(writer, sheet_name="Reposição", index=False)
        from openpyxl.styles import Font, PatternFill
        for ws in writer.book:
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions
            for cel in ws[1]:
                cel.font = Font(bold=True, color="FFFFFF")
                cel.fill = PatternFill("solid", fgColor="EB5E33")
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 28
                for cel in col[1:]:
                    if isinstance(cel.value, (int, float)):
                        cel.number_format = 'R$ #,##0.00' if ws.title == "Verba em estoque" else "#,##0"
        if reposicao is not None:
            ws = writer.book["Reposição"]
            for row in ws.iter_rows(min_row=2):
                estilos = destacar_reposicao(pd.Series([c.value for c in row], index=COLUNAS_REPOSICAO))
                for cel, estilo in zip(row, estilos):
                    if estilo:
                        cel.fill = PatternFill("solid", fgColor="713D13")
                        cel.font = Font(bold=True, color="FFFFFF")
        # O Excel também abre com os dois saldos; o sinal + revela os detalhes.
        writer.book["Estoque por unidade"].column_dimensions.group("D", "I", hidden=True)
    return buf.getvalue()


def pagina_estoque():
    # Além de ficar fora do menu público, protege leituras e downloads caso a
    # função seja chamada diretamente. O app revalida o prazo antes de rotear.
    if not st.session_state.get("lanc_ok"):
        st.error("Área restrita. Libere o acesso no cadeado da barra lateral.")
        return

    st.title("📦 Estoque")
    st.caption("Acompanhamento interno · entradas, vendas e posição de estoque por unidade.")
    try:
        vendas = db.ler_vendas_verbas()
        sell_out = db.ler_base_tidy()
        if vendas.empty or sell_out.empty:
            st.info("As bases de Sell In e Sell Out precisam estar preenchidas para acompanhar o estoque.")
            return
        # Valida classificação e tarifas com as mesmas regras da conciliação.
        mensal = calcular_verbas(vendas, sell_out, {}, {})
    except ValueError as exc:
        st.error(str(exc))
        return
    except Exception:
        st.error("Não foi possível carregar as bases de estoque. Verifique a conexão e tente novamente.")
        return

    ultima_in = pd.to_datetime(vendas["data"]).max()
    ultima_out = pd.to_datetime(sell_out["mes"]).max()
    st.caption(f"Sell In até {ultima_in:%d/%m/%Y} · Sell Out com lançamentos até {ultima_out:%m/%Y}. O último mês pode estar parcial.")
    if ultima_out.to_period("M") < ultima_in.to_period("M"):
        st.warning("O Sell Out termina antes do Sell In. O estoque pode incluir produtos cujas vendas ainda não foram lançadas.")

    ca, cm = st.columns([1, 2])
    ano = ca.selectbox("Ano", sorted(set(mensal.index.year), reverse=True), key="estoque_ano")
    meses = list(reversed(mensal[mensal.index.year == ano].index))
    mes = cm.selectbox("Posição ao final de", meses,
                       format_func=lambda m: f"{MESES[m.month]}/{m.year}", key="estoque_mes")
    movimento = movimentacao_por_unidade(vendas, sell_out, mes)
    unidades = st.multiselect("Unidades", list(movimento.index),
                              placeholder="Todas as unidades", key="estoque_unidades")
    selecionadas = unidades or list(movimento.index)
    movimento = movimento.loc[selecionadas]
    totais = movimento.sum()
    reposicao = reposicao_por_unidade(vendas, sell_out, mes)
    reposicao = reposicao[reposicao["Unidade"].isin(selecionadas)]
    alertas = reposicao[reposicao["Oferecer pedido"]]
    unidades_repor = alertas["Unidade"].nunique()

    historico = pd.DataFrame([
        movimentacao_por_unidade(vendas, sell_out, m).reindex(selecionadas, fill_value=0).sum()
        for m in mensal.index if m.year == ano and m <= mes
    ], index=[m for m in mensal.index if m.year == ano and m <= mes])
    historico.index.name = "Mês"
    divergencias = movimento[(movimento[["estoque_diant", "estoque_tras"]] < 0).any(axis=1)]

    cards = st.columns(4)
    cards[0].metric("Dianteiros em estoque (pares)", int(totais.estoque_diant))
    cards[1].metric("Traseiros em estoque (un)", int(totais.estoque_tras))
    cards[2].metric("Unidades com saldo negativo", len(divergencias))
    cards[3].metric("Unidades para novo pedido", unidades_repor)
    if not divergencias.empty:
        st.warning("Há saídas superiores às entradas registradas. Confira estoque inicial, transferências e cobertura das bases nas unidades destacadas abaixo.")
        with st.expander("Unidades para conferência", expanded=True):
            st.dataframe(divergencias[["estoque_diant", "estoque_tras"]].rename(columns=COLUNAS_ESTOQUE),
                         width="stretch")
    elif (historico[["estoque_diant", "estoque_tras"]] < 0).any().any():
        st.warning("O histórico do recorte contém saldo negativo. Confira as entradas e saídas nos meses anteriores.")

    st.subheader("Oportunidades de reposição")
    if unidades_repor:
        st.info(f"Oferecer novo pedido a {unidades_repor} unidade(s): o estoque de pelo menos um tipo está abaixo do consumo médio mensal.")
    else:
        st.info("Nenhuma unidade com consumo registrado está abaixo da média mensal no recorte selecionado.")
    reposicao_exibida = tabela_reposicao(reposicao)
    numeros = {c: "{:.0f}" for c in COLUNAS_REPOSICAO if c.startswith(("Saldo", "Consumo"))}
    st.dataframe(reposicao_exibida.style.apply(destacar_reposicao, axis=1).format(numeros, na_rep="—"),
                 hide_index=True, width="stretch")
    st.caption("Todas as unidades do filtro aparecem acima. O destaque indica qual produto oferecer: dianteiros, traseiros ou ambos. Dianteiros em pares; traseiros em unidades.")
    st.caption("Duração estimada: d = dias; m = meses de 30 dias. Exemplo: 1m e 12d = 42 dias de estoque. O consumo mensal é exibido sem casas decimais; a duração e os alertas usam a média original.")
    with st.expander("Como calculamos o consumo médio"):
        st.caption("Vendas totais ao consumidor ÷ meses com registros da unidade, desde o início do histórico até a posição selecionada, incluindo anos anteriores. Meses registrados com zero entram na média; meses sem registros não entram. O mês mais recente pode estar parcial.")

    st.subheader("Movimentação por unidade")
    exibicao = movimento.copy()
    exibicao.loc["TOTAL"] = totais
    saldos = exibicao[["estoque_diant", "estoque_tras"]].rename(columns={
        "estoque_diant": "Saldo Dianteiro", "estoque_tras": "Saldo Traseiro"})
    st.dataframe(saldos.style.format("{:.0f}"), width="stretch")
    st.caption("Dianteiros em pares; traseiros em unidades.")
    with st.expander("Ver estoque inicial, entradas e saídas", expanded=False):
        st.caption("Estoque inicial + entradas do mês − saídas do mês = saldo final.")
        detalhes = exibicao[["inicial_diant", "entrada_diant", "saida_diant",
                             "inicial_tras", "entrada_tras", "saida_tras"]].rename(columns=COLUNAS_ESTOQUE)
        st.dataframe(detalhes.style.format("{:.0f}"), width="stretch")

    st.subheader("Evolução do estoque")
    grafico = historico[["estoque_diant", "estoque_tras"]].rename(columns={
        "estoque_diant": "Dianteiros (pares)", "estoque_tras": "Traseiros (un)"})
    grafico.index = [MESES[m.month][:3] for m in grafico.index]
    fig = px.line(grafico, markers=True, color_discrete_sequence=["#EB5E33", "#F5A623"])
    fig.update_layout(template="plotly_dark", xaxis_title=None, yaxis_title="Quantidade",
                      legend_title=None, paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)")
    st.plotly_chart(fig, width="stretch")
    hist_exibicao = historico.copy()
    hist_exibicao.index = [f"{MESES[m.month]}/{m.year}" for m in historico.index]
    hist_exibicao.index.name = "Mês"
    with st.expander("Entradas e saídas mês a mês"):
        st.dataframe(hist_exibicao.rename(columns=COLUNAS_ESTOQUE).style.format("{:,.0f}"), width="stretch")

    valores = {}
    for cat, campo in CATEGORIAS.items():
        valores[cat] = sum(
            float(vendas.loc[vendas.tipo_refil == tipo, f"verba_{campo}"].iloc[0]) * totais[f"estoque_{tipo}"]
            for tipo in ("diant", "tras") if (vendas.tipo_refil == tipo).any()
        )
    verba = pd.DataFrame([{"Categoria": cat.title(), "Verba em estoque": round(valor, 2)}
                          for cat, valor in valores.items()])
    verba.loc[len(verba)] = ["Total", round(sum(valores.values()), 2)]
    with st.expander("Verba associada ao estoque"):
        st.dataframe(verba.style.format({"Verba em estoque": dinheiro}), hide_index=True, width="stretch")
        st.caption("Quantidade em estoque × tarifa por tipo de refil. Esta verba já integra o saldo geral de Verbas; não deve ser somada novamente. Os valores acompanham as unidades selecionadas.")
    with st.expander("Critérios de conferência"):
        st.markdown(
            "O estoque acumula todo o histórico disponível, com saldo inicial zero no início das bases. "
            "O filtro de ano não reinicia o saldo. Entradas vêm do Sell In e saídas, do Sell Out.\n\n"
            "Saldos negativos são mantidos para identificar entradas faltantes ou transferências entre lojas. "
            "Jeep Interlagos corresponde a Jeep W. Luiz."
        )
    st.download_button("Exportar acompanhamento de estoque (.xlsx)",
                       exportar_estoque(movimento, hist_exibicao, verba, reposicao),
                       file_name=f"estoque_{mes:%Y_%m}.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
